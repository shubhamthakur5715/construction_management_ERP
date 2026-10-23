from decimal import Decimal
from django.contrib import messages
from django.contrib.auth.decorators import login_required
from django.db.models import Sum
from django.shortcuts import render, redirect, get_object_or_404
from django.http import HttpResponse
from django.utils import timezone
from django.core.exceptions import PermissionDenied
from .models import Project, Worker, Material, MaterialTransaction, Expense, Payment, LabourAttendance, Invoice, UserProfile
from .forms import ProjectForm, WorkerForm, MaterialForm, MaterialTransactionForm, ExpenseForm, PaymentForm, LabourAttendanceForm, InvoiceForm, UserProfileForm
from .permissions import role_required, get_role, visible_projects

def project_filter(request, qs):
    project_id = request.GET.get("project", "")
    if project_id and project_id.isdigit(): qs = qs.filter(project_id=project_id)
    return qs, project_id

@login_required
def dashboard(request):
    projects = visible_projects(request.user)
    expenses = Expense.objects.filter(project__in=projects)
    payments = Payment.objects.filter(project__in=projects)
    workers = Worker.objects.filter(project__in=projects, active=True)
    materials = Material.objects.filter(project__in=projects)
    context = {
        "projects": projects.order_by("-created_at")[:6], "project_count": projects.count(),
        "active_count": projects.filter(status="In Progress").count(), "worker_count": workers.count(),
        "expense_total": expenses.aggregate(total=Sum("amount"))["total"] or 0,
        "payment_total": payments.aggregate(total=Sum("amount"))["total"] or 0,
        "material_count": materials.count(), "low_stock": materials.filter(quantity__lte=0).count(),
        "role": get_role(request.user),
    }
    return render(request, "projects/dashboard.html", context)

@login_required
def project_list(request):
    qs = visible_projects(request.user).order_by("-created_at")
    status = request.GET.get("status", "")
    if status: qs = qs.filter(status=status)
    return render(request, "projects/project_list.html", {"projects": qs, "selected_status": status})

@login_required
@role_required("Admin", "Site Manager")
def project_create(request):
    form = ProjectForm(request.POST or None)
    if form.is_valid():
        form.save(); messages.success(request, "Project created successfully."); return redirect("project_list")
    return render(request, "projects/form.html", {"form": form, "heading": "Add Project"})

@login_required
@role_required("Admin", "Site Manager")
def project_edit(request, pk):
    obj = get_object_or_404(visible_projects(request.user), pk=pk)
    form = ProjectForm(request.POST or None, instance=obj)
    if form.is_valid():
        form.save(); messages.success(request, "Project updated."); return redirect("project_list")
    return render(request, "projects/form.html", {"form": form, "heading": "Edit Project"})

@login_required
@role_required("Admin")
def project_delete(request, pk):
    obj = get_object_or_404(Project, pk=pk)
    if request.method == "POST":
        obj.delete(); messages.success(request, "Project deleted."); return redirect("project_list")
    return render(request, "projects/confirm_delete.html", {"object": obj})

def list_create(model, form_class, template_name, heading, url_name, allowed=("Admin","Accountant","Site Manager"), project_field=True):
    @login_required
    @role_required(*allowed)
    def listing(request):
        objects = model.objects.all().order_by("-pk")
        if project_field and hasattr(model, "project"): objects = objects.filter(project__in=visible_projects(request.user))
        objects, selected = project_filter(request, objects) if project_field else (objects, "")
        return render(request, template_name, {"objects": objects, "heading": heading, "url_name": url_name, "create_url": url_name.replace("_list", "_create"), "projects": visible_projects(request.user), "selected_project": selected})
    @login_required
    @role_required(*allowed)
    def create(request):
        form = form_class(request.POST or None)
        if hasattr(form, "fields") and "project" in form.fields:
            form.fields["project"].queryset = visible_projects(request.user)
        if form.is_valid():
            obj = form.save()
            if hasattr(obj, "project") and obj.project_id and not visible_projects(request.user).filter(pk=obj.project_id).exists():
                obj.delete(); raise PermissionDenied("Project access denied.")
            messages.success(request, f"{heading.rstrip('s')} saved successfully.")
            return redirect(url_name)
        return render(request, "projects/form.html", {"form": form, "heading": "Add " + heading.rstrip("s")})
    return listing, create

worker_list, worker_create = list_create(Worker, WorkerForm, "projects/simple_list.html", "Workers", "worker_list")
material_list, material_create = list_create(Material, MaterialForm, "projects/simple_list.html", "Materials", "material_list")
expense_list, expense_create = list_create(Expense, ExpenseForm, "projects/simple_list.html", "Expenses", "expense_list")
payment_list, payment_create = list_create(Payment, PaymentForm, "projects/simple_list.html", "Payments", "payment_list", ("Admin","Accountant"))
attendance_list, attendance_create = list_create(LabourAttendance, LabourAttendanceForm, "projects/simple_list.html", "Attendance", "attendance_list")
invoice_list, invoice_create = list_create(Invoice, InvoiceForm, "projects/simple_list.html", "Invoices", "invoice_list", ("Admin","Accountant"))

@login_required
@role_required("Admin","Site Manager")
def material_transaction_create(request):
    form = MaterialTransactionForm(request.POST or None)
    form.fields["material"].queryset = Material.objects.filter(project__in=visible_projects(request.user))
    if form.is_valid():
        tx = form.save(commit=False)
        material = tx.material
        # Usage must not create negative stock; purchases and adjustments are handled by the model.
        if tx.transaction_type == "Usage" and tx.quantity > material.quantity:
            form.add_error("quantity", "Usage quantity cannot exceed current stock.")
        else:
            tx.save()
            if tx.transaction_type == "Purchase":
                Material.objects.filter(pk=material.pk).update(unit_cost=tx.unit_cost or material.unit_cost, supplier_name=tx.supplier_name or material.supplier_name)
            messages.success(request, "Stock transaction saved.")
            return redirect("material_list")
    return render(request, "projects/form.html", {"form": form, "heading": "Record Material Purchase / Usage"})

@login_required
@role_required("Admin")
def user_roles(request):
    users = UserProfile.objects.select_related("user").prefetch_related("projects").all()
    return render(request, "projects/user_roles.html", {"profiles": users})

@login_required
@role_required("Admin")
def user_role_edit(request, pk):
    profile = get_object_or_404(UserProfile, pk=pk)
    form = UserProfileForm(request.POST or None, instance=profile)
    if form.is_valid():
        form.save(); messages.success(request, "Role and project access updated."); return redirect("user_roles")
    return render(request, "projects/form.html", {"form": form, "heading": "Edit User Permissions"})

@login_required
def profit_loss(request):
    projects = visible_projects(request.user)
    selected = request.GET.get("project", "")
    if selected.isdigit(): projects = projects.filter(pk=selected)
    rows = []
    for p in projects:
        expenses = p.expenses.aggregate(v=Sum("amount"))["v"] or Decimal("0")
        wages = sum((a.wage_amount for a in p.attendance_records.select_related("worker")), Decimal("0"))
        material_cost = sum((m.stock_value for m in p.materials.all()), Decimal("0"))
        payments = p.payments.aggregate(v=Sum("amount"))["v"] or Decimal("0")
        invoiced = sum((i.total_amount for i in p.invoices.all()), Decimal("0"))
        rows.append({"project":p,"expenses":expenses,"wages":wages,"material_cost":material_cost,"payments":payments,"invoiced":invoiced,"budget":p.budget,"cost":expenses+wages,"variance":p.budget-expenses-wages})
    return render(request, "projects/profit_loss.html", {"rows":rows,"projects":visible_projects(request.user),"selected_project":selected})

@login_required
def invoice_detail(request, pk):
    invoice = get_object_or_404(Invoice.objects.filter(project__in=visible_projects(request.user)), pk=pk)
    return render(request, "projects/invoice_detail.html", {"invoice":invoice})

@login_required
@role_required("Admin","Accountant")
def export_report(request, kind):
    if kind not in ("projects","expenses","workers"): return redirect("dashboard")
    projects = visible_projects(request.user)
    project_id = request.GET.get("project", "")
    if project_id.isdigit(): projects = projects.filter(pk=project_id)
    if kind == "projects":
        headers = ["Project","Client","Site","Budget","Status","Progress"]
        rows = [[p.name,p.client_name,p.site_address,str(p.budget),p.status,f"{p.progress}%"] for p in projects]
        filename="project-summary"
    elif kind == "expenses":
        headers=["Date","Project","Category","Description","Amount"]
        data=Expense.objects.filter(project__in=projects).order_by("-date")
        rows=[[str(x.date),x.project.name,x.category,x.description,str(x.amount)] for x in data]; filename="expenses"
    else:
        headers=["Worker","Phone","Role","Project","Daily wage","Active"]
        data=Worker.objects.filter(project__in=projects).select_related("project")
        rows=[[x.name,x.phone,x.role,x.project.name if x.project else "",str(x.daily_wage),"Yes" if x.active else "No"] for x in data]; filename="workers"
    if request.GET.get("format") == "xlsx":
        from openpyxl import Workbook

        wb=Workbook(); ws=wb.active; ws.title=kind.title(); ws.append(headers)
        for row in rows: ws.append(row)
        for cell in ws[1]:
            cell.font = __import__("openpyxl").styles.Font(bold=True, color="FFFFFF")
            cell.fill = __import__("openpyxl").styles.PatternFill("solid", fgColor="263449")
        for col in ws.columns:
            width=max(len(str(c.value or "")) for c in col)+2
            ws.column_dimensions[col[0].column_letter].width=min(width,45)
        response=HttpResponse(content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        response["Content-Disposition"]=f'attachment; filename="{filename}.xlsx"'; wb.save(response); return response
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
    from reportlab.lib import colors
    from reportlab.lib.styles import getSampleStyleSheet

    response=HttpResponse(content_type="application/pdf"); response["Content-Disposition"]=f'attachment; filename="{filename}.pdf"'
    doc=SimpleDocTemplate(response,pagesize=landscape(A4),rightMargin=24,leftMargin=24,topMargin=28,bottomMargin=28)
    styles=getSampleStyleSheet(); story=[Paragraph(f"Construction Report: {kind.title()}",styles["Title"]),Spacer(1,12)]
    data=[headers]+rows; table=Table(data,repeatRows=1)
    table.setStyle(TableStyle([("BACKGROUND",(0,0),(-1,0),colors.HexColor("#263449")),("TEXTCOLOR",(0,0),(-1,0),colors.white),("FONTNAME",(0,0),(-1,0),"Helvetica-Bold"),("FONTSIZE",(0,0),(-1,-1),8),("GRID",(0,0),(-1,-1),.35,colors.HexColor("#DDE3EA")),("ROWBACKGROUNDS",(0,1),(-1,-1),[colors.white,colors.HexColor("#F4F6F8")]),("VALIGN",(0,0),(-1,-1),"TOP"),("LEFTPADDING",(0,0),(-1,-1),6),("RIGHTPADDING",(0,0),(-1,-1),6)]))
    story.append(table); doc.build(story); return response
